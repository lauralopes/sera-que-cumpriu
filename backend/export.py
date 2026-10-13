import csv
import io
from openpyxl import Workbook
from .research import STATES

COLUMNS=[('id','ID'),('president','Presidente'),('election','Eleição'),('period','Período avaliado'),('theme','Tema'),('description','Promessa'),('quote','Trecho original'),('page','Página física PDF'),('target','Meta'),('deadline','Prazo'),('criterion','Critério de conclusão'),('status','Estado'),('result','Resultado observado'),('quantitative_progress','Avanço quantitativo'),('justification','Justificativa'),('sources','Fontes'),('confidence','Confiança'),('cutoff','Data de corte'),('human_review','Revisão humana'),('research_state','Pesquisa'),('criterion_note','Limitação do critério')]

def flatten(row):
    r=dict(row)
    r.update(period=f'{r.get("start","")} a {r.get("end","")}',quote=' | '.join(x['quote'] for x in r.get('refs',[])),page=', '.join(str(x['page']) for x in r.get('refs',[])),sources=' | '.join(x['url'] for x in r.get('evidence',[])))
    r['status']=r.get('status') or 'Pendente de pesquisa';r['human_review']='Sim' if r.get('human_review') else 'Não';return r

def csv_safe(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r','\n')):return "'"+value
    return value

def write_sheet(wb,name,headers,rows):
    ws=wb.create_sheet(name);ws.append(headers)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cells in ws.columns:ws.column_dimensions[cells[0].column_letter].width=min(55,max(16,len(str(cells[0].value))+3))
    return ws

def export_bytes(rows,format,summaries=None):
    flat=[flatten(r) for r in rows]
    values=[[r.get(k) for k,_ in COLUMNS] for r in flat]
    if format=='csv':
        output=io.StringIO();writer=csv.writer(output);writer.writerow([h for _,h in COLUMNS]);writer.writerows([[csv_safe(v) for v in row] for row in values]);return output.getvalue().encode('utf-8-sig')
    if format!='xlsx':raise ValueError('Formato inválido')
    wb=Workbook();wb.remove(wb.active)
    write_sheet(wb,'Compromissos',[h for _,h in COLUMNS],values)
    write_sheet(wb,'Evidências',['Compromisso','Título','Instituição','URL','Publicação','Período início','Período fim','Consulta','Dado de suporte','Tipo','Contexto','Entrega posterior','Arquivo'],
        [[r['id'],e.get('title'),e.get('institution'),e.get('url'),e.get('published_at'),e.get('covered_start'),e.get('covered_end'),e.get('consulted_at'),e.get('support'),e.get('kind'),e.get('context'),bool(e.get('after_exercise')),e.get('archive_url')] for r in rows for e in r.get('evidence',[])])
    write_sheet(wb,'Histórico',['Compromisso','Revisão','Estado','Resultado','Justificativa','Data','Responsável','Motivo','Evidências IDs','Revisão anterior','Critério na avaliação'],
        [[r['id'],a.get('id'),a.get('status'),a.get('result'),a.get('justification'),a.get('created_at'),a.get('actor'),a.get('change_reason'),a.get('evidence_ids'),a.get('previous_id'),a.get('criterion_snapshot')] for r in rows for a in r.get('history',[])])
    write_sheet(wb,'Histórico critérios',['Compromisso','Anterior','Novo','Motivo','Responsável','Data'],
        [[r['id'],a.get('previous_criterion'),a.get('criterion'),a.get('reason'),a.get('actor'),a.get('created_at')] for r in rows for a in r.get('criterion_history',[])])
    if summaries is not None:
        write_sheet(wb,'Resumo',['Plano','Extração parcial','Total extraído','Mensuráveis','Concluídos','Determinado','% principal','% determinado','Pendentes','Páginas revisadas','Páginas total'],
            [[s['id'],s['partial_universe'],s['total'],s['measurable'],s['completed'],s['determined'],s['principal_percent'],s['determined_percent'],s['pending'],s['pages_reviewed'],s['pages_total']] for s in summaries])
    write_sheet(wb,'Metodologia',['Regra'],[[x] for x in [
        'Universo parcial: valores não representam todos os compromissos do plano enquanto a extração não estiver concluída.',
        'Principal = concluídos mensuráveis / todos os mensuráveis extraídos × 100. Peso igual por compromisso.',
        'Determinado = concluídos / (concluídos + parciais + não concluídos + revertidos) × 100.',
        'Ausência de evidência não comprova descumprimento. Parciais não recebem 50% automático.',
        'Cobertura de pesquisa = mensuráveis avaliados / mensuráveis extraídos. Cobertura documental é separada.',
        'Fonte estatística única não sustenta avaliação. Evidência complementar tem proveniência e limitações.',
        'Avaliações iniciais produzidas por Codex sem revisão humana. Critérios piloto não foram cegos às fontes.',
        'Dilma 2014: último dia completo de exercício 11/05/2016; 12/05 requer revisão por horário; Temer separado.',
        'Lula 2022: avaliação provisória. Sem percentual agregado por presidente.',
        'Resumo descreve o banco completo; aba Compromissos pode conter apenas o filtro exportado.']])
    output=io.BytesIO();wb.save(output);return output.getvalue()
